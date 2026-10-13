"""赛项一取证数据勘察脚本：扫描全部文件，输出结构与质量画像。"""
import sys, os, json, hashlib, collections
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE / "libs"))
DATA = BASE / "data" / "赛项一：竞赛数据"

import openpyxl


def sniff_encoding(raw: bytes):
    for enc in ("utf-8-sig", "utf-8", "gb18030", "big5", "cp932", "latin-1"):
        try:
            raw.decode(enc)
            return enc, True
        except Exception:
            continue
    return "undecodable", False


def txt_report(p: Path):
    raw = p.read_bytes()
    enc, ok = sniff_encoding(raw)
    text = raw.decode(enc, errors="replace") if ok else raw.decode("latin-1")
    lines = text.splitlines()
    delim = "\t" if "\t" in (lines[0] if lines else "") else ("," if "," in (lines[0] if lines else "") else None)
    ncol = len(lines[0].split(delim)) if delim and lines else 0
    sizes = {len(l.split(delim)) for l in lines[1:] if delim}
    return {
        "bytes": len(raw),
        "lines": len(lines),
        "md5": hashlib.md5(raw).hexdigest(),
        "enc": enc,
        "delim": delim,
        "header": lines[0].split(delim) if delim and lines else lines[:1],
        "ragged_col_counts": sorted(sizes)[:10],
        "sample": lines[1:3] if len(lines) > 2 else [],
    }


def xlsx_report(p: Path):
    info = {"bytes": p.stat().st_size, "sheets": []}
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception as e:
        info["error"] = f"{type(e).__name__}: {e}"
        return info
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        n = 0
        blank = 0
        for r in rows:
            n += 1
            if r is None or all(v in (None, "") for v in r):
                blank += 1
        info["sheets"].append({
            "name": ws.title,
            "dims": f"{ws.max_row}x{ws.max_column}",
            "data_rows": n,
            "all_blank_rows": blank,
            "header": [str(h) for h in header] if header else [],
        })
    wb.close()
    return info


def main():
    out = {"files": []}
    groups = collections.Counter()
    for p in sorted(DATA.rglob("*")):
        if not p.is_file():
            continue
        rel = str(p.relative_to(DATA))
        kind = p.suffix.lower()
        rec = {"path": rel, "kind": kind}
        if kind == ".txt":
            rec.update(txt_report(p))
        elif kind == ".xlsx":
            rec.update(xlsx_report(p))
        else:
            rec["note"] = "unknown type"
        out["files"].append(rec)
        groups[rel.split("/")[0] if "/" in rel else rel] += 1
    out["group_counts"] = dict(groups)
    (BASE / "work" / "profile.json").parent.mkdir(exist_ok=True, parents=True)
    (BASE / "work" / "profile.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")

    # console summary
    print("GROUPS:", json.dumps(out["group_counts"], ensure_ascii=False))
    txts = [f for f in out["files"] if f["kind"] == ".txt"]
    print("TXT encodings:", collections.Counter(f.get("enc") for f in txts))
    print("TXT names:", collections.Counter(Path(f["path"]).name for f in txts))
    hdrs = collections.Counter(tuple(f.get("header") or []) for f in txts)
    print("distinct TXT headers:", len(hdrs))
    for h, c in hdrs.most_common(8):
        print("  x%-4d %s" % (c, " | ".join(h)[:300]))
    xlsx = [f for f in out["files"] if f["kind"] == ".xlsx"]
    print("XLSX count:", len(xlsx), "errors:", sum(1 for f in xlsx if "error" in f))
    sh = collections.Counter()
    for f in xlsx:
        for s in f.get("sheets", []):
            sh[tuple(s["header"])] += 1
    print("distinct XLSX headers:", len(sh))
    for h, c in sh.most_common(12):
        print("  x%-4d %s" % (c, " | ".join(h)[:400]))
    for f in xlsx[:3]:
        print("SAMPLE", f["path"], json.dumps(f, ensure_ascii=False)[:600])


if __name__ == "__main__":
    main()
